import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SOURCES = list(ROOT.glob("BD_MADRE_DE_COSTOS*.xlsx"))
if not SOURCES:
    raise FileNotFoundError(f"No se encontró BD_MADRE_DE_COSTOS*.xlsx en {ROOT}")
INPUT = max(SOURCES, key=lambda path: path.stat().st_mtime)
OUTPUT = Path(__file__).resolve().parents[1] / "dist" / "costs.json"
MONTHS = ["Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
ALFARO = "ALFARO"


def clean(value):
    if value is None:
        return ""
    return " ".join(str(value).replace("\ufffd", "").split())


def num(value):
    return round(float(value), 6) if isinstance(value, (int, float)) else None


def latest_cost(row):
    return next((item["value"] for item in reversed(row["costs"]) if item["value"] is not None), None)


def main():
    formulas = load_workbook(INPUT, data_only=False, read_only=False)
    values = load_workbook(INPUT, data_only=True, read_only=False)

    detail_ws = values.worksheets[3]
    alfaro_needs_correction = any(
        clean(row[3]).upper() == ALFARO and clean(row[12]) == "+18%"
        for row in detail_ws.iter_rows(min_row=5, values_only=True)
        if row[2]
    )

    base_f = formulas["BD MADRE DE COSTOS"]
    base_v = values["BD MADRE DE COSTOS"]
    base = []
    for r in range(5, base_v.max_row + 1):
        gtin = clean(base_v.cell(r, 2).value)
        product = clean(base_v.cell(r, 3).value)
        provider = clean(base_v.cell(r, 4).value)
        if not (product or provider):
            continue
        costs = []
        for c in range(5, 13):
            value = num(base_v.cell(r, c).value)
            if provider.upper() == ALFARO and alfaro_needs_correction and value is not None:
                # ALFARO's invoice amount already includes IGV. The supplied workbook
                # added 18% again, so remove only that duplicated tax adjustment.
                value = round(value / 1.18, 4)
            formula = base_f.cell(r, c).value
            costs.append({"value": value, "carried": isinstance(formula, str) and formula.startswith("=")})
        base.append({
            "n": base_v.cell(r, 1).value,
            "gtin": gtin,
            "product": product,
            "provider": provider,
            "costs": costs,
            "providerCount": base_v.cell(r, 13).value or 0,
            "unitsPerBox": num(base_v.cell(r, 14).value),
            "unitsPurchased": num(base_v.cell(r, 15).value),
            "monthsPurchased": base_v.cell(r, 16).value or 0,
        })

    base_by_key = {}
    for item in base:
        base_by_key.setdefault(item["gtin"] or item["product"], []).append(item)

    # Recalculate the comparison from the corrected base. Otherwise ALFARO could
    # retain an obsolete minimum/maximum and an incorrect estimated saving.
    ws = values["COMPARATIVO PROVEEDORES"]
    comparisons = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        gtin = clean(row[1])
        product = clean(row[2])
        if not product or clean(row[0]).upper() == "TOTAL":
            continue
        candidates = []
        for item in base_by_key.get(gtin or product, []):
            value = latest_cost(item)
            if value is not None:
                candidates.append((value, item["provider"]))
        if not candidates:
            continue
        cheapest_cost, cheapest = min(candidates, key=lambda item: item[0])
        expensive_cost, expensive = max(candidates, key=lambda item: item[0])
        difference = round(expensive_cost - cheapest_cost, 4)
        units_purchased = num(row[11]) or 0
        comparisons.append({
            "n": 0, "gtin": gtin, "product": product,
            "unitsPerBox": num(row[3]), "providerCount": row[4] or 0,
            "cheapest": cheapest, "cheapestCost": cheapest_cost,
            "mostExpensive": expensive, "mostExpensiveCost": expensive_cost,
            "difference": difference,
            "differencePct": round(difference / cheapest_cost * 100, 1) if cheapest_cost else None,
            "unitsPurchased": units_purchased,
            "estimatedSaving": round(difference * units_purchased, 2),
        })
    comparisons.sort(key=lambda item: item["estimatedSaving"] or 0, reverse=True)
    for index, item in enumerate(comparisons, 1):
        item["n"] = index

    ws = values["PRESENTACIONES DISTINTAS"]
    presentations = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[2]:
            presentations.append({"n": row[0], "gtin": clean(row[1]), "product": clean(row[2]), "detail": clean(row[3])})

    ws = detail_ws
    detail = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row[2]:
            continue
        provider = clean(row[3])
        cost_without_tax = num(row[11])
        tax = clean(row[12])
        final_cost = num(row[13])
        if provider.upper() == ALFARO:
            tax = "IGV incluido"
            final_cost = cost_without_tax
        detail.append({
            "n": row[0], "gtin": clean(row[1]), "description": clean(row[2]), "provider": clean(row[3]),
            "month": clean(row[4]), "invoice": clean(row[5]), "file": clean(row[6]),
            "boxes": num(row[7]), "unitsPerBox": num(row[8]), "units": num(row[9]),
            "invoiceAmount": num(row[10]), "costWithoutTax": cost_without_tax, "tax": tax,
            "finalCost": final_cost, "amountSource": clean(row[14]), "unitsSource": clean(row[15]),
        })

    payload = {
        "meta": {
            "source": INPUT.name,
            "months": MONTHS,
            "baseRows": len(base),
            "products": len({x["gtin"] or x["product"] for x in base}),
            "gtinCount": len({x["gtin"] for x in base if x["gtin"]}),
            "comparisons": len(comparisons),
            "differentPresentations": len(presentations),
            "detailRows": len(detail),
            "igvIncludedProviders": ["ALFARO", "CASTILLO", "FP MAYORISTA (QUIMICA SUIZA)", "FDS"],
            "alfaroCorrectionApplied": alfaro_needs_correction,
        },
        "base": base,
        "comparisons": comparisons,
        "presentations": presentations,
        "detail": detail,
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps(payload["meta"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
