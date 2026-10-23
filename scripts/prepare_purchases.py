"""Build a safe purchase/provider dataset and conservative GTIN proposals."""
from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PARENT = ROOT.parent
PURCHASES = PARENT / "COMPRAS"
DETAIL_SOURCE = PURCHASES / "Consolidado Facturas" / "consolidado.xlsx"
REGISTER_SOURCE = PURCHASES / "COMPRAS AVANTI FARMA 2026.xlsx"
CATALOG_SOURCE = PURCHASES / "BDD DISTRIBUIDOR - LABORATORIO por medicamento.xlsx"
SALES_JSON = ROOT / "dist" / "sales.json"
OUTPUT = ROOT / "dist" / "purchases.json"

PROVIDER_RULES = {
    "ALFARO": {"display": "ALFARO", "included": True, "aliases": ["ALFARO"]},
    "CASTILLO": {"display": "CASTILLO", "included": True, "aliases": ["CASTILLO"]},
    "EBENEL": {"display": "EBENEL", "included": True, "aliases": ["EBENEL"]},
    "FDS": {"display": "FDS", "included": True, "aliases": ["FARMACEUTICA DEL SUR", "FDS"]},
    "FP MAYORISTA (QUIMICA SUIZA)": {"display": "FP MAYORISTA (QUIMICA SUIZA)", "included": True, "aliases": ["FP MAYORISTA", "QUIMICA SUIZA"]},
    "DIMEXA": {"display": "DIMEXA", "included": False, "aliases": ["DIMEXA"]},
    "MYM": {"display": "MYM", "included": False, "aliases": ["M&M", "MYM", "PROMEFAR"]},
    "SANTA URSULA": {"display": "SANTA URSULA", "included": False, "aliases": ["SANTA URSULA"]},
    "ALCON": {"display": "ALCON", "included": False, "aliases": ["ALCON"]},
    "PIERINELLA": {"display": "PIERINELLA", "included": False, "aliases": ["ANGYMARC", "PIERINELLA"]},
    "FARMAGOLD": {"display": "FARMAGOLD", "included": False, "aliases": ["FARMAGOLD"]},
}

GENERIC = {
    "CAJA", "CAJ", "CJA", "UND", "UNIDAD", "UNIDADES", "BIENES", "PRODUCTO",
    "COMP", "COMPR", "COMPRIMIDO", "COMPRIMIDOS", "TAB", "TABLETA", "TABLETAS",
    "CAP", "CAPS", "CAPSULA", "CAPSULAS", "REC", "RECUB", "FCO", "FRASCO",
    "SOL", "SOLUCION", "ORAL", "INY", "INYECTABLE", "GTS", "GOTAS", "DSCTO",
    "BONIF", "BONIFICACION", "N", "NUEVO", "NVA", "NVO",
}


def normalize(value) -> str:
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn").upper()
    return re.sub(r"\s+", " ", text).strip()


def provider_name(value) -> str:
    raw = normalize(value).replace(".", "")
    for key, rule in PROVIDER_RULES.items():
        if any(normalize(alias).replace(".", "") in raw for alias in rule["aliases"]):
            return key
    cleaned = re.sub(r"\b(SAC|S A C|SA|S A|SRL|S R L|EIRL|E I R L)\b", "", raw)
    return re.sub(r"\s+", " ", cleaned).strip() or "PROVEEDOR SIN NOMBRE"


def number(value):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).upper().replace("S/", "").replace("PEN", "").replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(",", "") if text.rfind(".") > text.rfind(",") else text.replace(".", "").replace(",", ".")
    elif "," in text:
        tail = text.rsplit(",", 1)[-1]
        text = text.replace(",", ".") if len(tail) <= 2 else text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def iso_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    for pattern in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            pass
    return ""


def identifier(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)) and float(value).is_integer():
        return str(int(value))
    return str(value).strip().replace(".0", "")


def core_name(value) -> str:
    text = normalize(value).replace("®", " ").replace("×", " X ")
    text = re.sub(r"\([^)]*30(?:04|03)\.[^)]*\)", " ", text)
    text = re.sub(r"\[[^]]*\]", " ", text)
    text = re.sub(r"\b(?:DSCTO|BONIF(?:ICACION)?|LOTE|LOT)\b.*$", " ", text)
    text = re.sub(r"\s+-\s+[A-Z][A-Z .&/-]*$", " ", text)
    text = re.sub(r"(?<=\d),(?=\d)", ".", text)
    text = re.sub(r"[^A-Z0-9.%+]+", " ", text)
    replacements = {"CJA": "CAJA", "CAJ": "CAJA", "COMP": "TAB", "COMPR": "TAB", "COMPRIMIDOS": "TAB", "COMPRIMIDO": "TAB", "TABLETAS": "TAB", "TABLETA": "TAB", "CAPSULAS": "CAP", "CAPSULA": "CAP", "FRASCO": "FCO", "GOTAS": "GTS"}
    parts = [replacements.get(token, token) for token in text.split()]
    return " ".join(parts)


def match_tokens(value):
    return [token for token in core_name(value).split() if token not in GENERIC]


def brand(value) -> str:
    for token in match_tokens(value):
        if not re.fullmatch(r"\d+(?:\.\d+)?(?:MG|MCG|UG|G|ML|UI|%)?", token):
            return token
    return ""


def strengths(value):
    text = core_name(value).replace(" ", "")
    return set(re.findall(r"\d+(?:\.\d+)?(?:MG|MCG|UG|GR|G|ML|UI|%)", text))


def pack_count(value):
    text = core_name(value)
    hits = []
    for match in re.finditer(r"(?:\bX|CAJA\s*X|CAJA\s*\*)\s*(\d{1,4})(?![\d.,])", text):
        rest = text[match.end():].lstrip()
        if re.match(r"^(?:ML|MG|MCG|UG|GR|G|KG|L|CM|MM|UI|DOSIS)\b", rest):
            continue
        hits.append(int(match.group(1)))
    return hits[0] if len(set(hits)) == 1 else None


def score_names(left, right):
    a, b = core_name(left), core_name(right)
    at, bt = set(match_tokens(left)), set(match_tokens(right))
    jaccard = len(at & bt) / len(at | bt) if at | bt else 0
    return 0.55 * jaccard + 0.45 * SequenceMatcher(None, a, b).ratio()


def load_sales_names():
    if not SALES_JSON.exists():
        return {}
    payload = json.loads(SALES_JSON.read_text(encoding="utf-8"))
    names = defaultdict(Counter)
    for row in payload.get("rows", []):
        if row.get("code"):
            names[str(row["code"])][row.get("product", "")] += 1
    return {code: counts.most_common(1)[0][0] for code, counts in names.items()}


def load_catalog(sales_names):
    workbook = load_workbook(CATALOG_SOURCE, read_only=True, data_only=True)
    sheet = workbook["CONSOLIDADO"]
    entries = []
    for row in sheet.iter_rows(min_row=7, values_only=True):
        # The catalog starts in column B; column A is deliberately blank.
        gtin, name = identifier(row[4]), str(row[5] or "").strip()
        if not gtin or not name:
            continue
        entries.append({"provider": provider_name(row[1]), "gtin": gtin, "name": name,
                        "systemName": sales_names.get(gtin, ""), "core": core_name(name),
                        "brand": brand(name), "strengths": strengths(name), "pack": pack_count(name)})
    workbook.close()
    # The system sales code is the target GTIN. Add its own description as a
    # global reference, without assigning it to a purchase provider.
    for gtin, name in sales_names.items():
        if gtin and name:
            entries.append({"provider": "__SYSTEM__", "gtin": gtin, "name": name,
                            "systemName": name, "core": core_name(name), "brand": brand(name),
                            "strengths": strengths(name), "pack": pack_count(name)})
    deduped = {}
    for item in entries:
        deduped[(item["provider"], item["gtin"], item["core"])] = item
    return list(deduped.values())


def propose(description, provider, catalog):
    provider_pool = [item for item in catalog if item["provider"] == provider]
    core = core_name(description)
    exact = {(item["gtin"], item["name"], item["systemName"]) for item in provider_pool if item["core"] == core}
    if len(exact) == 1:
        gtin, name, system = exact.pop()
        return {"status": "Coincidencia exacta", "gtin": gtin, "catalogName": name, "systemName": system, "score": 1, "method": "normalized_exact"}
    global_exact = {(item["gtin"], item["name"], item["systemName"]) for item in catalog if item["core"] == core}
    if len({item[0] for item in global_exact}) == 1:
        gtin, name, system = sorted(global_exact)[0]
        return {"status": "Coincidencia exacta global", "gtin": gtin, "catalogName": name, "systemName": system, "score": 1, "method": "global_exact"}
    b = brand(description)
    same_brand = [item for item in catalog if b and item["brand"] == b]
    candidates = [item for item in provider_pool if item["brand"] == b]
    seen = {(item["gtin"], item["core"]) for item in candidates}
    candidates.extend(item for item in same_brand if (item["gtin"], item["core"]) not in seen)
    if not candidates:
        candidates = provider_pool
    if not candidates:
        return {"status": "Sin catálogo del proveedor", "gtin": "", "catalogName": "", "systemName": "", "score": None, "method": "none"}
    source_strengths, source_pack = strengths(description), pack_count(description)
    ranked = []
    for item in candidates:
        if source_strengths and item["strengths"] and source_strengths != item["strengths"]:
            continue
        if source_pack and item["pack"] and source_pack != item["pack"]:
            continue
        ranked.append((score_names(description, item["name"]), item))
    ranked.sort(key=lambda pair: pair[0], reverse=True)
    if not ranked:
        return {"status": "Sin coincidencia", "gtin": "", "catalogName": "", "systemName": "", "score": None, "method": "none"}
    top_score, top = ranked[0]
    second_score = ranked[1][0] if len(ranked) > 1 else 0
    if top_score >= 0.86 and top_score - second_score >= 0.08:
        status = "Propuesta alta"
    elif top_score >= 0.70:
        status = "Revisar propuesta"
    else:
        status = "Sin coincidencia"
    if status == "Sin coincidencia":
        return {"status": status, "gtin": "", "catalogName": "", "systemName": "", "score": round(top_score, 4), "method": "fuzzy"}
    return {"status": status, "gtin": top["gtin"], "catalogName": top["name"], "systemName": top["systemName"],
            "score": round(top_score, 4), "method": "fuzzy", "secondScore": round(second_score, 4)}


def load_detail(catalog):
    workbook = load_workbook(DETAIL_SOURCE, read_only=True, data_only=True)
    rows, cache = [], {}
    for sheet in workbook.worksheets:
        headers = [str(v or "").strip() for v in next(sheet.iter_rows(values_only=True))]
        columns = {name: i for i, name in enumerate(headers)}
        for source_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            provider_raw = values[columns["Proveedor"]]
            description = str(values[columns["Descripcion del Producto"]] or "").strip()
            if not provider_raw or not description:
                continue
            provider = provider_name(provider_raw)
            key = (provider, core_name(description))
            if key not in cache:
                cache[key] = propose(description, provider, catalog)
            amount = number(values[columns["Importe Total Item"]])
            kind = str(values[columns["Tipo"]] or "").strip()
            signed_amount = -abs(amount) if amount is not None and "CREDITO" in normalize(kind) else amount
            rows.append({
                "sourceSheet": sheet.title, "sourceRow": source_row, "provider": provider,
                "providerOriginal": str(provider_raw).strip(), "date": iso_date(values[columns["Fecha de Emision"]]),
                "document": str(values[columns["Documento"]] or "").strip(), "type": kind,
                "qty": number(values[columns["Cantidad"]]), "unit": str(values[columns["Unidad de Medida"]] or "").strip(),
                "description": description, "lot": str(values[columns["Numero de Lote"]] or "").strip(),
                "unitPrice": number(values[columns["Precio Unitario"]]), "amount": signed_amount,
                "bonus": not signed_amount or "BONI" in normalize(description), "file": str(values[columns["Archivo"]] or "").strip(),
                "mappingKey": f"{provider}|{core_name(description)}", **cache[key],
            })
    workbook.close()
    mappings = [{"key": f"{provider}|{core}", "provider": provider, "description": next(r["description"] for r in rows if r["mappingKey"] == f"{provider}|{core}"), **mapping}
                for (provider, core), mapping in cache.items()]
    return rows, mappings


def load_register():
    workbook = load_workbook(REGISTER_SOURCE, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(v or "").strip() for v in next(iterator)]
    columns = {name: i for i, name in enumerate(headers) if name}
    rows = []
    for source_row, values in enumerate(iterator, 2):
        get = lambda name: values[columns[name]] if columns[name] < len(values) else None
        provider_raw = get("PROVEEDOR")
        amount = number(get("MONTO"))
        if not provider_raw or amount is None:
            continue
        kind = str(get(".") or "").strip()
        signed = -abs(amount) if "CREDITO" in normalize(kind) or normalize(kind).startswith("NC") else amount
        rows.append({"sourceRow": source_row, "provider": provider_name(provider_raw),
                     "providerOriginal": str(provider_raw).strip(), "document": str(get("NRO. DOC.") or "").strip(),
                     "date": iso_date(get("F. FACT.")), "amount": signed,
                     "status": str(get("STATUS") or "").strip(),
                     "deliveryDate": iso_date(get("F. ENTREGA")),
                     "dueDate": iso_date(get("F. VENC.")),
                     "paymentDate": iso_date(get("F. PAGO"))})
    workbook.close()
    return rows


def summarize(rows, register):
    groups = defaultdict(lambda: {"detailAmount": 0, "detailLines": 0, "documents": set(), "descriptions": set(),
                                  "registeredAmount": 0, "registeredRows": 0, "registeredDocuments": set(), "dates": []})
    for row in rows:
        group = groups[row["provider"]]
        group["detailAmount"] += row["amount"] or 0
        group["detailLines"] += 1
        group["documents"].add(row["document"])
        group["descriptions"].add(row["description"])
        if row["date"]:
            group["dates"].append(row["date"])
    for row in register:
        group = groups[row["provider"]]
        group["registeredAmount"] += row["amount"] or 0
        group["registeredRows"] += 1
        if row["document"]:
            group["registeredDocuments"].add(row["document"])
        if row["date"]:
            group["dates"].append(row["date"])
    providers = []
    for name, group in groups.items():
        rule = PROVIDER_RULES.get(name, {"included": False})
        providers.append({"provider": name, "igv": "Incluido" if rule["included"] else "Añadir 18%",
                          "detailAmount": round(group["detailAmount"], 2), "detailLines": group["detailLines"],
                          "detailDocuments": len(group["documents"]), "uniqueDescriptions": len(group["descriptions"]),
                          "registeredAmount": round(group["registeredAmount"], 2), "registeredRows": group["registeredRows"],
                          "registeredDocuments": len(group["registeredDocuments"]),
                          "from": min(group["dates"]) if group["dates"] else "", "to": max(group["dates"]) if group["dates"] else ""})
    return sorted(providers, key=lambda item: item["registeredAmount"], reverse=True)


def main():
    for source in (DETAIL_SOURCE, REGISTER_SOURCE, CATALOG_SOURCE):
        if not source.exists():
            raise FileNotFoundError(source)
    sales_names = load_sales_names()
    catalog = load_catalog(sales_names)
    detail, mappings = load_detail(catalog)
    register = load_register()
    providers = summarize(detail, register)
    mapping_counts = dict(Counter(item["status"] for item in mappings))
    dates = [row["date"] for row in detail + register if row["date"]]
    payload = {
        "meta": {"generated": date.today().isoformat(), "detailSource": DETAIL_SOURCE.name,
                 "registerSource": REGISTER_SOURCE.name, "catalogSource": CATALOG_SOURCE.name,
                 "from": min(dates), "to": max(dates), "detailRows": len(detail), "registerRows": len(register),
                 "providers": len(providers), "catalogRows": len(catalog), "uniqueMappings": len(mappings),
                 "mappingCounts": mapping_counts,
                 "privacy": "No contiene contactos, teléfonos, cuentas bancarias, credenciales ni observaciones del maestro de proveedores."},
        "rules": [{"provider": name, "igv": "Incluido" if rule["included"] else "Añadir 18%", "factor": 1 if rule["included"] else 1.18}
                  for name, rule in PROVIDER_RULES.items()],
        "providers": providers, "register": register, "rows": detail, "mappings": mappings,
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Compras: {len(detail):,} líneas de detalle; {len(register):,} registros de facturas; {len(providers)} proveedores")
    print("Conciliación:", ", ".join(f"{key}: {value}" for key, value in sorted(mapping_counts.items())))
    print(f"Generado: {OUTPUT}")


if __name__ == "__main__":
    main()
