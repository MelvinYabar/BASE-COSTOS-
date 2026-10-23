"""Prepare the daily sales report without publishing the private Excel source."""
import argparse
import json
from collections import Counter
from datetime import datetime, date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def number(value):
    if value is None or str(value).strip() == "":
        raise ValueError("Número ausente")
    return Decimal(str(value).replace("S/", "").replace(",", "").strip())


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", nargs="?", type=Path)
    args = parser.parse_args()
    candidates = list(ROOT.glob("*PRODUCTOS-VENDIDOS-POR-DIA*.xlsx"))
    if not args.source and not candidates:
        parser.error("Indica la ruta del Excel de ventas.")
    source = args.source or max(candidates, key=lambda p: p.stat().st_mtime)
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook["VENTAS-DIARIAS-POR-ITEM"]
    rows, qty_total, amount_total, reported = [], Decimal(0), Decimal(0), None
    header_found = False
    for n, cells in enumerate(sheet.iter_rows(values_only=True), 1):
        cells = list(cells) + [None] * 7
        if cells[1] == "Fecha" and cells[6] == "Total Vendido":
            header_found = True
            continue
        if not header_found:
            continue
        if str(cells[1]).strip().upper() == "TOTALES:":
            reported = (number(cells[5]), number(cells[6]))
            continue
        if all(v is None for v in cells[:7]):
            continue
        try:
            day = cells[1]
            day = day.date() if isinstance(day, datetime) else day
            if not isinstance(day, date):
                day = datetime.strptime(str(day), "%d/%m/%Y").date()
            code = "" if cells[2] is None else str(cells[2]).strip()
            if isinstance(cells[2], (int, float)):
                if int(cells[2]) != cells[2] or abs(cells[2]) > 2**53 - 1:
                    raise ValueError("Código numérico inválido; conservar como texto")
                code = str(int(cells[2]))
            qty, amount = number(cells[5]), number(cells[6])
            if qty == 0 or not cells[3]:
                raise ValueError("Cantidad, importe o descripción inválidos")
            rows.append({"date": day.isoformat(), "code": code, "product": str(cells[3]).strip(),
                         "qty": float(qty), "amount": float(amount), "provider": "Sin distribuidor",
                         "unit": "UNIDAD VENDIDA", "cost": None, "document": "", "sourceRow": n,
                         "type": str(cells[4] or ""), "signIssue": qty * amount < 0})
            qty_total += qty
            amount_total += amount
        except Exception as exc:
            raise ValueError(f"Fila {n}: {exc}") from exc
    workbook.close()
    if not rows or reported != (qty_total, amount_total):
        raise ValueError(f"No cuadra el reporte: observado {(qty_total, amount_total)}; declarado {reported}")
    payload = {"meta": {"source": source.name, "sheet": sheet.title, "rowCount": len(rows),
        "from": min(r["date"] for r in rows), "to": max(r["date"] for r in rows),
        "qtyTotal": float(qty_total), "amountTotal": float(amount_total),
        "missingCode": sum(not r["code"] for r in rows), "zeroAmount": sum(r["amount"] == 0 for r in rows),
        "codeCount": len({r["code"] for r in rows if r["code"]}),
        "reportedQty": float(reported[0]), "reportedAmount": float(reported[1]),
        "signIssues": sum(r["signIssue"] for r in rows),
        "types": dict(Counter(r["type"] for r in rows)), "hasDistributor": False}, "rows": rows}
    output = ROOT / "dist" / "sales.json"
    output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Ventas conciliadas: {len(rows):,} filas; cantidad {qty_total:,.2f}; S/ {amount_total:,.2f}")
    print(f"Sin código: {payload['meta']['missingCode']}; importe cero: {payload['meta']['zeroAmount']}")
    print(f"Generado: {output}")


if __name__ == "__main__":
    main()
