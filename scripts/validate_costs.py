import argparse
import json
import math
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DATA = Path(__file__).resolve().parents[1] / "dist" / "costs.json"
SOURCES = list(ROOT.glob("BD_MADRE_DE_COSTOS*.xlsx"))


def close(a, b, tolerance):
    return a is not None and b is not None and math.isclose(a, b, abs_tol=tolerance)


def main():
    parser = argparse.ArgumentParser(description="Valida la BD de costos y, si está disponible, su Excel fuente.")
    parser.add_argument("--published-only", action="store_true", help="Validar solo el JSON, sin comparación con el Excel original")
    args = parser.parse_args()
    errors = []
    warnings = []
    payload = json.loads(DATA.read_text(encoding="utf-8"))
    meta = payload["meta"]
    igv_included_providers = {provider.upper() for provider in meta.get("igvIncludedProviders", [])}

    expected_lengths = {
        "baseRows": len(payload["base"]),
        "comparisons": len(payload["comparisons"]),
        "differentPresentations": len(payload["presentations"]),
        "detailRows": len(payload["detail"]),
    }
    for field, observed in expected_lengths.items():
        if meta.get(field) != observed:
            errors.append(f"{field}: meta={meta.get(field)}; datos={observed}")

    base_keys = [(row.get("gtin") or row.get("product"), row.get("provider")) for row in payload["base"]]
    duplicates = [key for key, count in Counter(base_keys).items() if count > 1]
    if duplicates:
        warnings.append(f"Claves producto-proveedor repetidas para revisión: {len(duplicates)} ({'; '.join(f'{product} / {provider}' for product, provider in duplicates[:5])})")

    providers_by_product = defaultdict(set)
    for row in payload["base"]:
        product_key = row.get("gtin") or row.get("product")
        providers_by_product[product_key].add(row.get("provider"))
        observed_months = sum(1 for item in row["costs"] if item.get("value") is not None and not item.get("carried"))
        if observed_months != row.get("monthsPurchased"):
            errors.append(f"Fila base {row.get('n')}: meses con compra={row.get('monthsPurchased')}; observados={observed_months}")
        seen_cost = False
        for item in row["costs"]:
            if item.get("carried") and not seen_cost:
                errors.append(f"Fila base {row.get('n')}: costo arrastrado sin costo anterior")
            if item.get("value") is not None:
                seen_cost = True

    for row in payload["base"]:
        product_key = row.get("gtin") or row.get("product")
        if row.get("providerCount") != len(providers_by_product[product_key]):
            errors.append(f"Fila base {row.get('n')}: cantidad de proveedores inconsistente")

    for row in payload["comparisons"]:
        cheapest = row.get("cheapestCost")
        expensive = row.get("mostExpensiveCost")
        difference = row.get("difference")
        percentage = row.get("differencePct")
        if None not in (cheapest, expensive) and cheapest > expensive:
            errors.append(f"Comparativo {row.get('n')}: costo mínimo mayor al máximo")
        if None not in (cheapest, expensive, difference) and not close(expensive - cheapest, difference, 0.001):
            errors.append(f"Comparativo {row.get('n')}: diferencia monetaria inconsistente")
        if None not in (cheapest, expensive, percentage) and cheapest > 0:
            expected = (expensive - cheapest) / cheapest * 100
            if not close(expected, percentage, 0.6):
                errors.append(f"Comparativo {row.get('n')}: diferencia porcentual inconsistente")

        matching = [item for item in payload["base"] if (item.get("gtin") or item.get("product")) == (row.get("gtin") or row.get("product"))]
        candidates = [(next((cost.get("value") for cost in reversed(item["costs"]) if cost.get("value") is not None), None), item.get("provider")) for item in matching]
        candidates = [item for item in candidates if item[0] is not None]
        if candidates:
            expected_min = min(candidates, key=lambda item: item[0])
            expected_max = max(candidates, key=lambda item: item[0])
            if (cheapest, row.get("cheapest")) != expected_min:
                errors.append(f"Comparativo {row.get('n')}: proveedor/costo mínimo no coincide con la base corregida")
            if (expensive, row.get("mostExpensive")) != expected_max:
                errors.append(f"Comparativo {row.get('n')}: proveedor/costo máximo no coincide con la base corregida")

    for row in payload["detail"]:
        boxes = row.get("boxes")
        units_per_box = row.get("unitsPerBox")
        units = row.get("units")
        if None not in (boxes, units_per_box, units) and not close(boxes * units_per_box, units, 0.01):
            errors.append(f"Detalle {row.get('n')}: cajas × unidades/caja no coincide con unidades")
        without_tax = row.get("costWithoutTax")
        final_cost = row.get("finalCost")
        if row.get("provider", "").upper() in igv_included_providers:
            if row.get("tax") != "IGV incluido":
                errors.append(f"Detalle {row.get('n')}: {row.get('provider')} debe marcar IGV incluido")
            if None not in (without_tax, final_cost) and not close(without_tax, final_cost, 0.000001):
                errors.append(f"Detalle {row.get('n')}: {row.get('provider')} no debe recibir un segundo IGV")
        elif row.get("tax") == "+18%" and None not in (without_tax, final_cost) and not close(without_tax * 1.18, final_cost, 0.0015):
            errors.append(f"Detalle {row.get('n')}: costo final no coincide con IGV de 18%")

    missing_gtin = sum(1 for row in payload["detail"] if not row.get("gtin"))
    assumed_units = sum(1 for row in payload["detail"] if "se asume 1" in row.get("unitsSource", ""))
    warnings.extend([
        f"Filas de detalle sin GTIN: {missing_gtin}",
        f"Filas con una unidad por envase asumida: {assumed_units}",
    ])

    if args.published_only:
        warnings.append("Modo JSON: no se comparó con el Excel fuente original.")
    elif not SOURCES:
        errors.append(f"No se encontró BD_MADRE_DE_COSTOS*.xlsx en {ROOT}")
    else:
        source = max(SOURCES, key=lambda path: path.stat().st_mtime)
        workbook = load_workbook(source, data_only=True, read_only=True)
        source_counts = {
            "baseRows": sum(1 for row in workbook["BD MADRE DE COSTOS"].iter_rows(min_row=5, values_only=True) if row[2] or row[3]),
            "comparisons": sum(1 for row in workbook["COMPARATIVO PROVEEDORES"].iter_rows(min_row=5, values_only=True) if row[2] and str(row[0]).strip().upper() != "TOTAL"),
            "differentPresentations": sum(1 for row in workbook["PRESENTACIONES DISTINTAS"].iter_rows(min_row=5, values_only=True) if row[2]),
            "detailRows": sum(1 for row in workbook.worksheets[3].iter_rows(min_row=5, values_only=True) if row[2]),
        }
        for field, observed in source_counts.items():
            if observed != expected_lengths[field]:
                errors.append(f"{field}: Excel={observed}; JSON={expected_lengths[field]}")
        if meta.get("source") != source.name:
            errors.append(f"Fuente declarada={meta.get('source')}; fuente usada={source.name}")

    print(f"Errores críticos: {len(errors)}")
    for item in errors[:30]:
        print(f"ERROR: {item}")
    print(f"Advertencias de revisión humana: {len(warnings)}")
    for item in warnings:
        print(f"AVISO: {item}")
    if errors:
        raise SystemExit(1)
    print("VALIDACIÓN AUTOMÁTICA APROBADA")


if __name__ == "__main__":
    main()
